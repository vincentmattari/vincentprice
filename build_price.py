# -*- coding: utf-8 -*-
"""
재고 — 적정 판매가 산출 엔진
============================
입력: *_발주_입고_횟수_*.xlsx
      필수 — 부품 품번, 발주, 입고, CNT, 발주평균가, 입고평균가, 전체평균가, 전체최소가, 전체최대가
      선택 — 부품명, 메이커 (있으면 결과에 함께 싣는다)
      컬럼은 헤더 이름으로 찾으므로 순서가 바뀌거나 추가돼도 동작한다.
출력: index.html (페이지), data.json, 적정판매가_산출결과.xlsx

핵심 전제 (데이터로 검증됨)
---------------------------
1) 발주평균가 + 입고평균가 == 전체평균가  (2179/2179 행에서 성립)
   -> 두 컬럼은 '단가'가 아니라 CNT로 나눈 '기여분'이다.
   -> 실제 단가 복원:  발주단가 = 발주평균가 x CNT / 발주
2) 가격 0 레코드는 '무료'가 아니라 '미기재'다.
   근거: CNT=2 & 최소가=0 & 최대가>0 인 365행 전부가 평균가 == 최대가/2 를 정확히 만족.
        (한 건은 가격이 있고 한 건은 0 -> 평균이 정확히 절반으로 희석)
   -> 전체평균가는 구조적으로 과소평가되어 있다.
3) 전사 '가격 미기재' 비율 p0: CNT=1 표본 959건 중 410건이 가격 0 -> p0 = 0.428
"""
import openpyxl, json, math, sys, io, glob, statistics as st
from collections import Counter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

PATTERN = "*_발주_입고_횟수_*.xlsx"     # 원자료는 저장소에 포함하지 않으므로 패턴으로 찾는다
_found = sorted(glob.glob(PATTERN))
if not _found:
    sys.exit(f"원본 엑셀을 찾을 수 없습니다. 이 폴더에 {PATTERN} 파일을 두세요.")
SRC = _found[-1]
P0 = 0.45          # 가격 미기재 비율 사전값 (실측 0.428 + 여유)
MARGIN = {"보수": 1.15, "권장": 1.30, "공격": 1.45}


def round_price(v):
    """판매가 단수 정리: 1만원 미만 100원, 10만원 미만 500원, 그 이상 1000원 단위 올림"""
    if v <= 0:
        return 0
    unit = 100 if v < 10_000 else (500 if v < 100_000 else 1_000)
    return int(math.ceil(v / unit) * unit)


def grade_of(L, U, cnt, priced):
    if not priced:
        return "F", "가격정보 전무 — 산출 불가"
    r = L / U if U > 0 else 0
    if r >= 0.95:
        return ("A", "단일 관측가 (교차검증 불가)") if cnt == 1 else ("A", "가격 일관 — 확정 수준")
    if r >= 0.60:
        return "B", "경미한 편차 — 신뢰 가능"
    if r >= 0.30:
        return "C", "편차 큼 — 검토 권장"
    return "D", "편차 심각 — 수동 확인 필수"


def load(path):
    """헤더 이름으로 컬럼을 찾는다 — 컬럼이 추가되거나 순서가 바뀌어도 동작한다."""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    hdr = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    # 같은 뜻으로 쓰이는 헤더명들 — 첫 번째로 발견되는 것을 쓴다
    ALIAS = {
        "부품 품번": ["부품 품번", "부품품번", "품번"],
        "CNT":      ["CNT", "발주/입고 합계", "발주/입고합계", "합계"],
        "부품명":    ["부품명", "품명"],
        "메이커":    ["메이커", "제조사"],
    }
    need = ["부품 품번", "발주", "입고", "CNT", "발주평균가", "입고평균가",
            "전체평균가", "전체최소가", "전체최대가"]
    ix = {}
    for k in need + ["부품명", "메이커"]:
        ix[k] = next((hdr.index(a) for a in ALIAS.get(k, [k]) if a in hdr), None)
    missing = [k for k in need if ix[k] is None]
    if missing:
        sys.exit(f"필수 컬럼 누락: {missing}\n  발견된 헤더: {hdr}")
    return ix, list(ws.iter_rows(min_row=2, values_only=True))


def txt(v):
    """빈 문자열·공백·None 을 None 으로 정규화하고 앞뒤 공백을 제거한다."""
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def main():
    ix, rows = load(SRC)
    has_meta = ix["부품명"] is not None or ix["메이커"] is not None

    out = []
    for row in rows:
        g_ = lambda k: row[ix[k]]
        pn = str(g_("부품 품번")).strip()
        name = txt(row[ix["부품명"]]) if ix["부품명"] is not None else None
        maker = txt(row[ix["메이커"]]) if ix["메이커"] is not None else None
        po, rc, cnt = g_("발주"), g_("입고"), g_("CNT")
        pavg, ravg, avg = g_("발주평균가"), g_("입고평균가"), g_("전체평균가")
        vmin, vmax = g_("전체최소가"), g_("전체최대가")
        u_po = (pavg * cnt / po) if po > 0 else None      # 복원한 발주 단가
        u_rc = (ravg * cnt / rc) if rc > 0 else None      # 복원한 입고 단가

        priced = vmax > 0
        if not priced:
            L = U = C = 0.0
            z_hat = float(cnt)
        else:
            U = float(vmax)
            # 하한: 0 오염은 아래로만 작용하므로, 후보 중 최댓값이 최선의 하한
            L = max(float(avg), u_po or 0.0, u_rc or 0.0, 0.0)
            L = min(L, U)
            # 0(미기재) 보정: 관측 상한과 전사 사전값 중 작은 쪽
            z_gap = cnt * (1 - avg / U)                   # 데이터가 허용하는 최대 미기재 건수
            z_hat = max(0.0, min(z_gap, cnt * P0))
            n_eff = cnt - z_hat
            C = (avg * cnt / n_eff) if n_eff > 0.5 else U
            C = min(max(C, L), U)                        # [L, U] 로 클램프

        g, note = grade_of(L, U, cnt, priced)
        rec = {
            "품번": pn, "부품명": name, "메이커": maker,
            "발주": po, "입고": rc, "CNT": cnt,
            "전체평균가": int(avg), "전체최소가": int(vmin), "전체최대가": int(vmax),
            "발주단가": int(round(u_po)) if u_po is not None else None,
            "입고단가": int(round(u_rc)) if u_rc is not None else None,
            "미기재추정": round(z_hat, 1),
            "원가하한": int(round(L)), "유효원가": int(round(C)), "원가상한": int(round(U)),
            "등급": g, "비고": note,
        }
        for k, m in MARGIN.items():
            rec[k + "가"] = round_price(C * m) if priced else 0
        out.append(rec)

    out.sort(key=lambda r: (-r["CNT"], -r["유효원가"]))

    # ---- 요약 ----
    gc = Counter(r["등급"] for r in out)
    ok = [r for r in out if r["등급"] != "F"]
    lift = [r["유효원가"] / r["전체평균가"] for r in ok if r["전체평균가"] > 0]
    makers = Counter(r["메이커"] for r in out if r["메이커"])
    summary = {
        "총부품수": len(out),
        "산출가능": len(ok),
        "산출불가": gc["F"],
        "등급분포": {k: gc[k] for k in "ABCDF"},
        "평균가대비_상향률_중앙값": round(st.median(lift), 4) if lift else 0,
        "유효원가_합계": sum(r["유효원가"] for r in out),
        "전체평균가_합계": sum(r["전체평균가"] for r in out),
        "권장가_합계": sum(r["권장가"] for r in out),
        "마진": MARGIN, "미기재비율_사전값": P0,
        "메이커수": len(makers),
        "메이커목록": [m for m, _ in makers.most_common()],
        "부품명결측": sum(1 for r in out if not r["부품명"]),
    }

    with open("data.json", "w", encoding="utf-8") as f:
        json.dump({"summary": summary, "items": out}, f, ensure_ascii=False, separators=(",", ":"))

    # ---- 웹 페이지 (데이터 내장, 단일 파일) ----
    FIELDS = ["품번", "부품명", "메이커", "발주", "입고", "CNT",
              "전체평균가", "전체최소가", "전체최대가",
              "발주단가", "입고단가", "원가하한", "유효원가", "원가상한", "등급"]
    web = {
        "summary": {**summary, "기준일": "2026-08-12"},
        "items": [[r[k] for k in FIELDS] for r in out],
    }
    with open("template.html", encoding="utf-8") as f:
        html = f.read()
    html = html.replace("/*__DATA__*/",
                        json.dumps(web, ensure_ascii=False, separators=(",", ":")))
    with open("index.html", "w", encoding="utf-8") as f:
        f.write(html)

    # ---- 산출방법론.md -> 산출방법론.html ----
    import render_md
    with open("산출방법론.md", encoding="utf-8") as f:
        body = render_md.render(f.read())
    with open("template_method.html", encoding="utf-8") as f:
        shell = f.read()
    with open("산출방법론.html", "w", encoding="utf-8") as f:
        f.write(shell.replace("<!--__CONTENT__-->", body))

    # ---- 엑셀 출력 ----
    wb2 = openpyxl.Workbook()
    ws = wb2.active
    ws.title = "적정판매가"
    cols = ["품번", "부품명", "메이커", "발주", "입고", "CNT",
            "전체평균가", "전체최소가", "전체최대가",
            "발주단가", "입고단가", "미기재추정", "원가하한", "유효원가", "원가상한",
            "등급", "보수가", "권장가", "공격가", "비고"]
    ws.append(cols)
    for c in ws[1]:
        c.font = openpyxl.styles.Font(bold=True)
        c.fill = openpyxl.styles.PatternFill("solid", fgColor="1F2937")
        c.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
    for r in out:
        ws.append([r.get(c) for c in cols])
    ws.freeze_panes = "A2"
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(10, len(c) * 2)
    wb2.save("적정판매가_산출결과.xlsx")

    # ---- 콘솔 리포트 ----
    print("=" * 72)
    print("총 부품수 %d  |  산출가능 %d  |  산출불가 %d"
          % (summary["총부품수"], summary["산출가능"], summary["산출불가"]))
    print("등급분포:", summary["등급분포"])
    print("전체평균가 대비 유효원가 상향률 중앙값: %.1f%%" % ((summary["평균가대비_상향률_중앙값"] - 1) * 100))
    print("원가총액  전체평균가 기준 %d  ->  유효원가 기준 %d"
          % (summary["전체평균가_합계"], summary["유효원가_합계"]))
    if has_meta:
        print("메이커 %d개  |  부품명 결측 %d건" % (summary["메이커수"], summary["부품명결측"]))
    print("=" * 72)
    print("\n[검증] 발주1건 유가 + 입고1건 0원 유형 (진값 = 최대가 여야 함)")
    for r in out:
        if r["품번"] in ("51747255413", "5K0955978A", "5NA807568CYE4"):
            print("  %-16s %-14.14s 평균가%9d  유효원가%9d  최대가%9d  등급%s"
                  % (r["품번"], r["부품명"] or "-", r["전체평균가"],
                     r["유효원가"], r["전체최대가"], r["등급"]))
    print("\n[상위 CNT 부품]")
    for r in out[:8]:
        print("  %-16s %-4.4s %-16.16s CNT%4d 평균가%8d 유효%8d [%s] 권장%9d"
              % (r["품번"], r["메이커"] or "-", r["부품명"] or "-", r["CNT"],
                 r["전체평균가"], r["유효원가"], r["등급"], r["권장가"]))


if __name__ == "__main__":
    main()
